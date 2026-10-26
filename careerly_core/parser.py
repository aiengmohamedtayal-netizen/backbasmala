"""Step 1 - get text out of an uploaded CV (PDF / DOCX)."""
import io
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024  # same limit the frontend enforces


class CVError(ValueError):
    """Base class: the message is safe to show to the user (HTTP 422)."""


class UnsupportedFormat(CVError): ...
class FileTooLarge(CVError): ...
class EmptyDocument(CVError): ...


def extract_text(data: bytes, filename: str) -> str:
    if len(data) > MAX_BYTES:
        raise FileTooLarge("File is larger than 5 MB. Upload a smaller file.")

    ext = Path(filename).suffix.lower()
    # Check magic bytes too: never trust the extension alone.
    if ext == ".pdf" and data[:5] == b"%PDF-":
        reader = _pdf
    elif ext == ".docx" and data[:2] == b"PK":
        reader = _docx
    elif ext == ".doc":
        raise UnsupportedFormat("Old .doc files are not supported. Please save as .docx or PDF.")
    else:
        raise UnsupportedFormat("Upload a valid PDF or DOCX file.")

    try:
        text = reader(data)
    except Exception:  # library errors must never reach the user as a 500
        raise CVError("We could not read this file. It may be corrupted or password-protected.")

    if len(text.strip()) < 30:
        raise EmptyDocument("No readable text found. The file may be a scanned image.")
    return text


def _pdf(data: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


def _docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
